"""Native PPTX helpers. Geometry uses integer coordinates; text sizes are integer pt."""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx_xml_checks import normalize

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def set_gradient(shape, stops, angle=0):
    if len(stops) < 2 or [s[0] for s in stops] != sorted(s[0] for s in stops):
        raise ValueError('At least two sorted gradient stops required')
    for pos, color, alpha in stops:
        if not all(isinstance(v, int) and 0 <= v <= 100000 for v in (pos, alpha)):
            raise ValueError('Position and alpha must be integers in 0..100000')
        RGBColor.from_string(color)
    sp = shape._element.spPr
    for tag in ('solidFill', 'gradFill', 'noFill', 'pattFill', 'blipFill', 'grpFill'):
        for el in list(sp.findall(f'{{{A}}}{tag}')):
            sp.remove(el)
    grad = OxmlElement('a:gradFill'); grad.set('rotWithShape', '1')
    gslist = OxmlElement('a:gsLst')
    for pos, color, alpha in stops:
        gs = OxmlElement('a:gs'); gs.set('pos', str(pos))
        cl = OxmlElement('a:srgbClr'); cl.set('val', color)
        al = OxmlElement('a:alpha'); al.set('val', str(alpha)); cl.append(al)
        gs.append(cl); gslist.append(gs)
    grad.append(gslist)
    lin = OxmlElement('a:lin'); lin.set('ang', str(int(angle))); lin.set('scaled', '1')
    grad.append(lin); sp.append(grad); normalize(shape._element)


def set_polygon(shape, points, width, height):
    if len(points) < 3 or width <= 0 or height <= 0:
        raise ValueError('Nonempty polygon with positive size required')
    if not all(isinstance(v, int) for v in [width, height] + [v for p in points for v in p]):
        raise ValueError('OOXML geometry requires integer coordinates')
    sp = shape._element.spPr
    for tag in ('prstGeom', 'custGeom'):
        for el in list(sp.findall(f'{{{A}}}{tag}')):
            sp.remove(el)
    geom = OxmlElement('a:custGeom')
    for tag in ('avLst', 'gdLst', 'ahLst', 'cxnLst'):
        geom.append(OxmlElement('a:' + tag))
    re = OxmlElement('a:rect')
    for k, v in [('l','0'), ('t','0'), ('r','r'), ('b','b')]:
        re.set(k, v)
    geom.append(re)
    paths = OxmlElement('a:pathLst'); path = OxmlElement('a:path')
    path.set('w', str(width)); path.set('h', str(height))
    for i, (x, y) in enumerate(points):
        cmd = OxmlElement('a:moveTo' if i == 0 else 'a:lnTo')
        pt = OxmlElement('a:pt'); pt.set('x', str(x)); pt.set('y', str(y)); cmd.append(pt)
        path.append(cmd)
    path.append(OxmlElement('a:close')); paths.append(path); geom.append(paths)
    sp.append(geom); normalize(shape._element)


def add_text_runs(slide, name, runs, box, source_size=(1600,900), slide_size=(13.333333,7.5), font='宋体'):
    if not all(isinstance(size, int) and size > 0 for _, size, _ in runs):
        raise ValueError('Font sizes must be positive integer points')
    x, y, w, h = box
    def qx(v): return Inches(v / source_size[0] * slide_size[0])
    def qy(v): return Inches(v / source_size[1] * slide_size[1])
    sh = slide.shapes.add_textbox(qx(x), qy(y), qx(w), qy(h)); sh.name = name
    tf = sh.text_frame; tf.clear(); tf.word_wrap = False; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; para = tf.paragraphs[0]
    for text, size, color in runs:
        for i, piece in enumerate(text.split('\n')):
            if i: para = tf.add_paragraph()
            r = para.add_run(); r.text = piece; r.font.name = font
            r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(color)
            rp = r._r.get_or_add_rPr()
            for tag in ('latin', 'ea', 'cs'):
                el = rp.find(f'{{{A}}}{tag}')
                if el is None: el = OxmlElement('a:' + tag); rp.append(el)
                el.set('typeface', font)
    normalize(sh._element)
    return sh
